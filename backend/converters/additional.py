from __future__ import annotations

import base64
import html
import io
import zipfile
from pathlib import Path

import bleach
import fitz
import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from PIL import Image

SAFE_TAGS = ["a", "b", "blockquote", "br", "code", "del", "em", "h1", "h2", "h3", "h4", "h5", "h6", "hr", "i", "li", "ol", "p", "pre", "strong", "table", "tbody", "td", "th", "thead", "tr", "ul"]
SAFE_ATTRIBUTES = {"a": ["href", "title"]}


def sanitize_html(value: str) -> str:
    soup = BeautifulSoup(value, "html.parser")
    for node in soup.find_all(["script", "style", "iframe", "object", "embed"]): node.decompose()
    value = str(soup)
    return bleach.clean(value, tags=SAFE_TAGS, attributes=SAFE_ATTRIBUTES, protocols=["http", "https", "mailto"], strip=True)


def markdown_to_html_bytes(data: bytes) -> bytes:
    rendered = markdown.markdown(data.decode("utf-8"), extensions=["extra", "sane_lists"])
    return sanitize_html(rendered).encode("utf-8")


def _html_to_pdf(value: str) -> bytes:
    output = io.BytesIO()
    writer = fitz.DocumentWriter(output)
    page = fitz.paper_rect("a4")
    content = page + (54, 54, -54, -54)
    try:
        story = fitz.Story(f"<html><body>{sanitize_html(value)}</body></html>", user_css="body{font-family:sans-serif;font-size:11pt;line-height:1.5}table{border-collapse:collapse;width:100%}td,th{border:1px solid #888;padding:5px}img{max-width:100%}")
        story.write(writer, lambda _number, _filled: (page, content, None))
    finally:
        writer.close()
    return output.getvalue()


def markdown_to_pdf(data: bytes) -> bytes:
    return _html_to_pdf(markdown_to_html_bytes(data).decode("utf-8"))


def html_to_pdf(data: bytes) -> bytes:
    return _html_to_pdf(data.decode("utf-8"))


def text_to_pdf(data: bytes) -> bytes:
    return _html_to_pdf(f"<pre>{html.escape(data.decode('utf-8'))}</pre>")


def _append_html(document: Document, value: str) -> None:
    soup = BeautifulSoup(sanitize_html(value), "html.parser")
    for node in soup.children:
        if isinstance(node, NavigableString):
            if node.strip(): document.add_paragraph(str(node).strip())
        elif isinstance(node, Tag):
            name = node.name.lower()
            if name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                document.add_heading(node.get_text(" ", strip=True), level=int(name[1]))
            elif name in {"ul", "ol"}:
                style = "List Bullet" if name == "ul" else "List Number"
                for item in node.find_all("li", recursive=False): document.add_paragraph(item.get_text(" ", strip=True), style=style)
            elif name == "table":
                rows = node.find_all("tr")
                width = max((len(row.find_all(["th", "td"], recursive=False)) for row in rows), default=0)
                if rows and width:
                    table = document.add_table(rows=len(rows), cols=width)
                    for row_index, row in enumerate(rows):
                        for column_index, cell in enumerate(row.find_all(["th", "td"], recursive=False)):
                            table.cell(row_index, column_index).text = cell.get_text(" ", strip=True)
            else:
                text = node.get_text(" ", strip=True)
                if text: document.add_paragraph(text)


def html_to_docx(data: bytes) -> bytes:
    document = Document(); _append_html(document, data.decode("utf-8")); output = io.BytesIO(); document.save(output); return output.getvalue()


def markdown_to_docx(data: bytes) -> bytes:
    return html_to_docx(markdown_to_html_bytes(data))


def text_to_docx(data: bytes) -> bytes:
    document = Document()
    for paragraph in data.decode("utf-8").split("\n\n"): document.add_paragraph(paragraph)
    output = io.BytesIO(); document.save(output); return output.getvalue()


def docx_to_text(data: bytes) -> bytes:
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    if not parts: raise ValueError("No readable content was found in this document.")
    return ("\n\n".join(parts) + "\n").encode("utf-8")


def docx_to_html(data: bytes) -> bytes:
    document = Document(io.BytesIO(data)); parts: list[str] = []
    for paragraph in document.paragraphs:
        text = html.escape(paragraph.text)
        style = (paragraph.style.name if paragraph.style else "").lower()
        if style.startswith("heading") and style[-1:].isdigit(): parts.append(f"<h{style[-1]}>{text}</h{style[-1]}>")
        elif text: parts.append(f"<p>{text}</p>")
    for table in document.tables:
        rows = ["<tr>" + "".join(f"<td>{html.escape(cell.text)}</td>" for cell in row.cells) + "</tr>" for row in table.rows]
        parts.append("<table>" + "".join(rows) + "</table>")
    if not parts: raise ValueError("No readable content was found in this document.")
    return sanitize_html("".join(parts)).encode("utf-8")


def image_to_pdf(data: bytes) -> bytes:
    with Image.open(io.BytesIO(data)) as image:
        rgba = image.convert("RGBA")
        frame = Image.new("RGB", image.size, "white")
        frame.paste(rgba, mask=rgba.getchannel("A"))
        output = io.BytesIO(); frame.save(output, format="PDF", resolution=150); return output.getvalue()


def pdf_pages_to_images(data: bytes, image_format: str) -> bytes:
    document = fitz.open(stream=data, filetype="pdf"); output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for number, page in enumerate(document, start=1):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            extension = "jpg" if image_format == "jpeg" else image_format
            archive.writestr(f"page-{number}.{extension}", pixmap.tobytes(image_format))
    document.close(); return output.getvalue()


def compress_pdf(data: bytes) -> bytes:
    source = fitz.open(stream=data, filetype="pdf"); output = io.BytesIO()
    source.save(output, garbage=4, deflate=True, clean=True); source.close(); return output.getvalue()


def merge_pdfs(sources: list[bytes]) -> bytes:
    output = fitz.open()
    try:
        for data in sources:
            source = fitz.open(stream=data, filetype="pdf")
            output.insert_pdf(source); source.close()
        return output.tobytes(garbage=4, deflate=True)
    finally:
        output.close()


def images_to_pdf(sources: list[bytes]) -> bytes:
    return merge_pdfs([image_to_pdf(data) for data in sources])


def split_pdf(data: bytes, options: dict[str, object]) -> bytes:
    value = str(options.get("page_range", "")).strip()
    if not value: raise ValueError("Enter a page range, for example 1-3 or 1,3,5.")
    source = fitz.open(stream=data, filetype="pdf"); output = fitz.open(); pages: list[int] = []
    try:
        for item in value.split(","):
            bounds = item.strip().split("-")
            if len(bounds) == 1: start = end = int(bounds[0])
            elif len(bounds) == 2: start, end = map(int, bounds)
            else: raise ValueError("Invalid page range.")
            if start < 1 or end < start or end > source.page_count: raise ValueError("The page range is outside this PDF.")
            pages.extend(range(start - 1, end))
        for page in pages: output.insert_pdf(source, from_page=page, to_page=page)
        return output.tobytes(garbage=4, deflate=True)
    except (TypeError, ValueError) as exc:
        raise ValueError("Choose a valid page range within the PDF.") from exc
    finally:
        source.close(); output.close()
