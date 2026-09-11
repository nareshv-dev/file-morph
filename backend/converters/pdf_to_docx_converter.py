from __future__ import annotations

from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Pt


# 3x PDF coordinates equals 216 DPI. This keeps small text crisp in Word while
# keeping the generated document a practical size for normal PDFs.
PDF_RENDER_SCALE = 3.0
PAGE_EDGE_TOLERANCE = 1.0


def _open_pdf(source: Path | bytes) -> fitz.Document:
    try:
        if isinstance(source, bytes):
            return fitz.open(stream=source, filetype="pdf")
        return fitz.open(source)
    except Exception as exc:
        raise ValueError("The PDF could not be processed.") from exc


def _configure_page(document: Document, width: float, height: float) -> None:
    section = document.sections[0]
    section.page_width = Pt(width)
    section.page_height = Pt(height)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)
    section.gutter = Pt(0)

    normal = document.styles["Normal"]
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)


def _add_pdf_page(
    document: Document,
    page: fitz.Page,
    output_width: float,
    output_height: float,
    page_index: int,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Pt(0)
    paragraph.paragraph_format.right_indent = Pt(0)
    if page_index:
        paragraph.paragraph_format.page_break_before = True

    # Mixed-size PDF pages are fitted to the first page's Word section without
    # cropping or stretching them.
    available_width = max(1.0, output_width - PAGE_EDGE_TOLERANCE)
    available_height = max(1.0, output_height - PAGE_EDGE_TOLERANCE)
    fit_scale = min(available_width / page.rect.width, available_height / page.rect.height)
    image_width = page.rect.width * fit_scale
    image_height = page.rect.height * fit_scale

    pixmap = page.get_pixmap(
        matrix=fitz.Matrix(PDF_RENDER_SCALE, PDF_RENDER_SCALE),
        alpha=False,
    )
    picture = paragraph.add_run().add_picture(
        BytesIO(pixmap.tobytes("png")),
        width=Pt(image_width),
        height=Pt(image_height),
    )
    picture._inline.docPr.set("name", f"PDF page {page_index + 1}")
    picture._inline.docPr.set(
        "descr",
        f"Layout-preserved image of PDF page {page_index + 1}",
    )


def convert_pdf_to_docx(source: Path | bytes) -> bytes:
    """Create a DOCX whose pages visually match the source PDF.

    PDF is fixed-layout, while normal Word paragraphs reflow. Rendering each
    PDF page into a full-page Word image reliably preserves text, icons, rules,
    images, columns, and spacing instead of rearranging the document.
    """
    pdf = _open_pdf(source)
    try:
        if pdf.needs_pass:
            raise ValueError("Password-protected PDFs are not supported.")
        if not pdf.page_count:
            raise ValueError("The PDF does not contain any pages.")

        first_page = pdf[0]
        output_width = first_page.rect.width
        output_height = first_page.rect.height
        document = Document()
        _configure_page(document, output_width, output_height)

        for page_index, page in enumerate(pdf):
            _add_pdf_page(
                document,
                page,
                output_width,
                output_height,
                page_index,
            )

        output = BytesIO()
        document.save(output)
        return output.getvalue()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("The PDF could not be converted to a Word document.") from exc
    finally:
        pdf.close()
