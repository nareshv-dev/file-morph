from __future__ import annotations

import base64
import io
import zipfile

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.converters.markdown_converter import markdown_table, normalize_markdown
from backend.main import app
from backend.utils.file_validator import MAX_FILE_SIZE

client = TestClient(app)


def make_pdf(with_image: bool = False) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Sample title", fontsize=22)
    page.insert_text((72, 110), "A faithful paragraph from the PDF.", fontsize=11)
    page.insert_text((72, 128), "XML-safe\x12control text", fontsize=11)
    if with_image:
        png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=")
        page.insert_image(fitz.Rect(72, 135, 110, 173), stream=png)
    data = document.tobytes()
    document.close()
    return data


def make_docx(with_image: bool = False) -> bytes:
    document = Document()
    document.add_heading("Sample title", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Bold words").bold = True
    paragraph.add_run(" and regular words.")
    document.add_paragraph("First item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "one"
    if with_image:
        png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=")
        document.add_picture(io.BytesIO(png))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_health_endpoint() -> None:
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"status": "healthy", "service": "FileMorph"}


def test_valid_pdf_conversion() -> None:
    response = client.post("/api/convert", files={"file": ("sample.pdf", make_pdf(), "application/pdf")})
    assert response.status_code == 200
    assert "Sample title" in response.json()["markdown"]
    assert "faithful paragraph" in response.json()["markdown"]


def test_valid_docx_conversion_with_image() -> None:
    response = client.post(
        "/api/convert",
        files={"file": ("sample.docx", make_docx(True), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    result = response.json()
    assert "# Sample title" in result["markdown"]
    assert "**Bold words**" in result["markdown"]
    assert "| Name | Value |" in result["markdown"]
    assert result["assets"]
    assert "![" in result["markdown"]
    assert result["bundle"]
    with zipfile.ZipFile(io.BytesIO(base64.b64decode(result["bundle"]))) as bundle:
        assert "sample.md" in bundle.namelist()
        assert any(name.startswith("images/") for name in bundle.namelist())


def test_pdf_to_docx_conversion() -> None:
    response = client.post(
        "/api/convert/pdf-to-docx",
        files={"file": ("sample.pdf", make_pdf(True), "application/pdf")},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert response.headers["x-output-filename"] == "sample.docx"
    converted = Document(io.BytesIO(response.content))
    assert len(converted.inline_shapes) == 1
    section = converted.sections[0]
    assert section.page_width.pt == pytest.approx(595.0, abs=1.0)
    assert section.page_height.pt == pytest.approx(842.0, abs=1.0)
    assert section.left_margin.pt == pytest.approx(0.0)
    assert section.top_margin.pt == pytest.approx(0.0)

    page_image = converted.inline_shapes[0]
    assert page_image.width.pt == pytest.approx(section.page_width.pt, abs=2.0)
    assert page_image.height.pt == pytest.approx(section.page_height.pt, abs=2.0)


def test_docx_to_pdf_conversion() -> None:
    response = client.post(
        "/api/convert/docx-to-pdf",
        files={"file": ("sample.docx", make_docx(True), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.headers["x-output-filename"] == "sample.pdf"
    converted = fitz.open(stream=response.content, filetype="pdf")
    text = "\n".join(page.get_text() for page in converted)
    converted.close()
    assert "Sample title" in text
    assert "Bold words" in text
    assert "alpha" in text


def test_conversion_rejects_wrong_source_type() -> None:
    response = client.post(
        "/api/convert/pdf-to-docx",
        files={"file": ("sample.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 400
    assert "PDF" in response.json()["error"]


@pytest.mark.parametrize(
    ("name", "data", "content_type", "expected"),
    [
        ("notes.txt", b"hello", "text/plain", "Only PDF and DOCX"),
        ("empty.pdf", b"", "application/pdf", "empty"),
        ("broken.pdf", b"not a pdf", "application/pdf", "invalid or corrupted"),
    ],
)
def test_invalid_uploads(name: str, data: bytes, content_type: str, expected: str) -> None:
    response = client.post("/api/convert", files={"file": (name, data, content_type)})
    assert response.status_code == 400
    assert expected.lower() in response.json()["error"].lower()


def test_oversized_file() -> None:
    response = client.post(
        "/api/convert",
        files={"file": ("huge.pdf", b"%PDF-" + b"0" * MAX_FILE_SIZE, "application/pdf")},
    )
    assert response.status_code == 413


def test_markdown_normalization() -> None:
    assert normalize_markdown(["# Title", "", "Paragraph\n\n\n\nNext"]) == "# Title\n\nParagraph\n\nNext\n"
    assert normalize_markdown(["Visible\x12 text"]) == "Visible text\n"
    assert markdown_table([["A", "B"], ["x", "y"]]) == "| A | B |\n| --- | --- |\n| x | y |"
