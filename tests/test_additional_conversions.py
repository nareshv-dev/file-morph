from __future__ import annotations

import io
import zipfile

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from backend.main import app

client = TestClient(app)


def source(extension: str) -> tuple[bytes, str]:
    if extension == "pdf":
        pdf = fitz.open(); page = pdf.new_page(); page.insert_text((72, 72), "Format sample"); output = pdf.tobytes(); pdf.close(); return output, "application/pdf"
    if extension == "docx":
        document = Document(); document.add_heading("Format sample", level=1); document.add_paragraph("Readable paragraph."); output = io.BytesIO(); document.save(output); return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if extension in {"png", "jpg", "webp"}:
        image = Image.new("RGB", (64, 64), "green"); output = io.BytesIO(); image.save(output, {"png": "PNG", "jpg": "JPEG", "webp": "WEBP"}[extension]); return output.getvalue(), {"png": "image/png", "jpg": "image/jpeg", "webp": "image/webp"}[extension]
    if extension == "md": return b"# Format sample\n\nReadable paragraph.", "text/markdown"
    if extension == "html": return b"<h1>Format sample</h1><p>Readable paragraph.</p><script>alert(1)</script>", "text/html"
    return b"Format sample\n\nReadable paragraph.", "text/plain"


@pytest.mark.parametrize("converter,extension,target", [
    ("markdown-to-pdf", "md", "pdf"),
    ("markdown-to-docx", "md", "docx"),
    ("markdown-to-html", "md", "html"),
    ("docx-to-html", "docx", "html"),
    ("docx-to-text", "docx", "txt"),
    ("html-to-pdf", "html", "pdf"),
    ("html-to-docx", "html", "docx"),
    ("text-to-pdf", "txt", "pdf"),
    ("text-to-docx", "txt", "docx"),
    ("image-to-pdf", "png", "pdf"),
    ("image-to-pdf", "jpg", "pdf"),
    ("image-to-pdf", "webp", "pdf"),
    ("pdf-to-png", "pdf", "zip"),
    ("pdf-to-jpg", "pdf", "zip"),
    ("compress-pdf", "pdf", "pdf"),
])
def test_additional_converter_with_real_file(converter: str, extension: str, target: str) -> None:
    data, mime = source(extension)
    response = client.post(f"/api/convert/{converter}", files={"file": (f"sample.{extension}", data, mime)})
    assert response.status_code == 200, response.text[:300]
    assert response.content
    if target == "pdf":
        pdf = fitz.open(stream=response.content, filetype="pdf"); assert pdf.page_count >= 1; pdf.close()
    elif target == "docx":
        document = Document(io.BytesIO(response.content)); assert any("Format sample" in p.text for p in document.paragraphs)
    elif target == "zip":
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive: assert archive.namelist()
    else:
        assert b"Format sample" in response.content
        assert b"<script" not in response.content
